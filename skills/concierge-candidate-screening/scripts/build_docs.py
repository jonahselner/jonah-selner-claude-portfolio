#!/usr/bin/env python3
"""Build the Executive Summary and Candidate Review .docx files for a screening batch.

Usage:
    python build_docs.py batch.json --outdir OUTDIR [--no-check]

batch.json format
-----------------
{
  "batch_date": "2026-10-26",                 # YYYY-MM-DD, used in header + filename
  "signoff": "Thanks for the strong batch ...", # one sentence; signed as Call Center Hiring Manager
  "intro": "...",                              # optional; a default two-sentence intro is used
  "candidates": [
    {
      "first_name": "Ana",
      "last_name": "Torres",
      "verdict": "GREEN" | "YELLOW" | "RED",
      "brief": "6 yrs Telvista inbound CS · 3.5-yr stint · native ES, strong EN",
      "partial": "resume only" | "transcript only" | null,
      "location": "Guadalajara, MX",
      "sources": "Resume, recruiter transcript, recruiter video",
      "missing": "Spanish not directly tested",   # or "" / null
      "why": ["...", "...", "..."],                 # <= 3 short bullets (Exec Summary)
      "checks": ["...", "..."],                     # YELLOW only, <= 3; mirrored in both docs
      "verdict_line": "One-line rationale",
      "rationale": "2-3 sentence paragraph",
      "positives": [{"c": 2, "text": "..."}],       # <= 4, c = criterion number 1-10
      "concerns":  [{"c": 4, "text": "..."}],       # <= 3
      "scorecard": {"1": "GREEN", "2": "GREEN", ..., "10": "N-A"}  # all 10 keys
    }
  ]
}

After writing, the script converts each file to PDF with LibreOffice (if available) and
reports any document that runs past one page so you can shorten the text and rerun.
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CRITERIA = {
    1: "Stable tenure",
    2: "2+ yrs phone CS",
    3: "English (written/spoken)",
    4: "Spanish fluency",
    5: "Empathy / EQ",
    6: "Coachability / professionalism",
    7: "Home office",
    8: "Shift availability",
    9: "Longevity mindset",
    10: "Resume / communication clarity",
}
COLORS = {  # fill, text
    "GREEN": ("D7F0DD", "1E6B34"),
    "YELLOW": ("FFF1C2", "8A6100"),
    "RED": ("F8D7D7", "9B1C1C"),
    "N-A": ("EEEEEE", "555555"),
}
DEFAULT_INTRO = (
    "GREEN candidates move straight to a personal interview with our team; RED candidates are a pass "
    "for this role. YELLOW candidates are close, and each has one specific check listed that will "
    "decide the call."
)
SIGNER = "Call Center Hiring Manager, Boost"


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def fix_widths(table, widths):
    """Set fixed column widths in a way both Word and LibreOffice respect."""
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.inches * 1440)))
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w


def cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def base_doc(font_size=9.5):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(font_size)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.05
    return doc


def para(container, text="", bold=False, size=None, color=None, italic=False, after=2, align=None):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def bullet(container, text, size=None, prefix="• "):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    r = p.add_run(prefix + text)
    if size:
        r.font.size = Pt(size)
    return p


def write_cell(cell, lines, size=8.5, first_bold=False):
    cell.text = ""
    first = cell.paragraphs[0]
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        r.font.size = Pt(size)
        if first_bold and i == 0:
            r.bold = True


def verdict_cell(cell, verdict, size=9):
    fill, txt = COLORS.get(verdict, COLORS["N-A"])
    shade(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(verdict)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(txt)


def full_name(c):
    return f"{c['first_name']} {c['last_name']}".strip()


def safe(s):
    # Strip accents (María -> Maria) so filenames stay readable, then collapse anything else.
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")


def build_summary(batch, outdir, scale=1.0):
    cands = batch["candidates"]
    n = len(cands)
    body = 8.5 * scale
    doc = base_doc(9.5 * scale)
    para(doc, "Candidate Screening — Executive Summary", bold=True, size=15 * scale, color="1F3A5F", after=1)
    para(doc, f"{batch['batch_date']} · {n} candidate{'s' if n != 1 else ''} reviewed · Boost",
         size=9.5 * scale, color="555555", after=4)
    para(doc, batch.get("intro") or DEFAULT_INTRO, size=9 * scale, after=5)

    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t)
    widths = [Inches(2.0), Inches(0.8), Inches(2.45), Inches(2.05)]
    for i, h in enumerate(["Candidate", "Decision", "Why", "Next step (YELLOW only)"]):
        c = t.rows[0].cells[i]
        shade(c, "1F3A5F")
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(8.5 * scale)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for cand in cands:
        row = t.add_row().cells
        name_lines = [full_name(cand)]
        brief = cand.get("brief", "")
        if cand.get("partial"):
            brief = f"{brief} — {cand['partial']}"
        name_lines.append(brief)
        write_cell(row[0], name_lines, size=body, first_bold=True)
        verdict_cell(row[1], cand["verdict"], size=9 * scale)
        write_cell(row[2], ["• " + w for w in cand.get("why", [])[:3]], size=body)
        if cand["verdict"] == "YELLOW" and cand.get("checks"):
            write_cell(row[3], ["• " + ch for ch in cand["checks"][:3]], size=body)
        else:
            write_cell(row[3], ["—"], size=body)
    fix_widths(t, widths)

    counts = {v: sum(1 for c in cands if c["verdict"] == v) for v in ("GREEN", "YELLOW", "RED")}
    para(doc, "", after=2)
    para(doc, f"{counts['GREEN']} GREEN · {counts['YELLOW']} YELLOW · {counts['RED']} RED",
         bold=True, size=10 * scale, after=4)
    p = para(doc, batch.get("signoff", ""), size=9 * scale, after=1)
    para(doc, f"— {SIGNER}", italic=True, size=9 * scale, color="555555")

    path = os.path.join(outdir, f"Executive_Summary_{batch['batch_date']}.docx")
    doc.save(path)
    return path


def build_review(cand, outdir, scale=1.0):
    doc = base_doc(9.5 * scale)
    v = cand["verdict"]
    fill, txt = COLORS[v]

    banner = doc.add_table(rows=1, cols=2)
    cell_margins(banner, 80, 80, 110, 110)
    left, right = banner.rows[0].cells
    shade(left, fill)
    shade(right, fill)
    left.text = ""
    r = left.paragraphs[0].add_run(full_name(cand))
    r.bold = True
    r.font.size = Pt(16 * scale)
    meta = left.add_paragraph()
    meta.paragraph_format.space_after = Pt(0)
    mr = meta.add_run(cand.get("location") or "Location not stated")
    mr.font.size = Pt(8.5 * scale)
    right.text = ""
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = rp.add_run(v)
    rr.bold = True
    rr.font.size = Pt(20 * scale)
    rr.font.color.rgb = RGBColor.from_string(txt)
    fix_widths(banner, [Inches(5.3), Inches(2.0)])

    src = f"Sources reviewed: {cand.get('sources', '—')}"
    if cand.get("missing"):
        src += f"   ·   Missing / not assessed: {cand['missing']}"
    para(doc, src, size=8 * scale, color="555555", after=4).paragraph_format.space_before = Pt(3)

    para(doc, cand.get("verdict_line", ""), bold=True, size=10.5 * scale, after=2)
    para(doc, cand.get("rationale", ""), size=9.5 * scale, after=5)

    def heading(text):
        para(doc, text, bold=True, size=10 * scale, color="1F3A5F", after=1).paragraph_format.space_before = Pt(3)

    if v == "YELLOW" and cand.get("checks"):
        heading("What the internal team needs to do to decide")
        for ch in cand["checks"][:3]:
            bullet(doc, ch, size=9.5 * scale, prefix="☐ ")
    heading("Positives")
    for item in cand.get("positives", [])[:4]:
        bullet(doc, f"[#{item['c']}] {item['text']}", size=9.5 * scale)
    if not cand.get("positives"):
        bullet(doc, "None noted", size=9.5 * scale)
    heading("Concerns / things to confirm")
    for item in cand.get("concerns", [])[:3]:
        bullet(doc, f"[#{item['c']}] {item['text']}", size=9.5 * scale)
    if not cand.get("concerns"):
        bullet(doc, "None noted", size=9.5 * scale)

    heading("Criterion scorecard")
    sc = cand.get("scorecard", {})
    t = doc.add_table(rows=5, cols=4)
    t.style = "Table Grid"
    cell_margins(t, 30, 30, 70, 70)
    for i in range(10):
        num = i + 1
        row, col = i % 5, (i // 5) * 2
        label, val = t.rows[row].cells[col], t.rows[row].cells[col + 1]
        write_cell(label, [f"{num}. {CRITERIA[num]}"], size=8.5 * scale)
        verdict_cell(val, str(sc.get(str(num), sc.get(num, "N-A"))).upper(), size=8 * scale)
    fix_widths(t, [Inches(2.75), Inches(0.9), Inches(2.75), Inches(0.9)])

    path = os.path.join(outdir, f"Candidate_Review_{safe(cand['first_name'])}_{safe(cand['last_name'])}.docx")
    doc.save(path)
    return path


def page_count(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as td:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, path],
                           check=True, capture_output=True, timeout=120)
        except Exception:
            return None
        pdf = os.path.join(td, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        if not os.path.exists(pdf):
            return None
        data = open(pdf, "rb").read()
        return len(re.findall(rb"/Type\s*/Page[^s]", data))


def validate(batch):
    errs = []
    for c in batch.get("candidates", []):
        n = full_name(c)
        if c.get("verdict") not in ("GREEN", "YELLOW", "RED"):
            errs.append(f"{n}: verdict must be GREEN/YELLOW/RED")
        if c.get("verdict") == "YELLOW" and not c.get("checks"):
            errs.append(f"{n}: YELLOW needs at least one specific check")
        if len(c.get("why", [])) > 3:
            errs.append(f"{n}: max 3 'why' bullets")
        if len(c.get("positives", [])) > 4 or len(c.get("concerns", [])) > 3:
            errs.append(f"{n}: max 4 positives / 3 concerns")
        missing = [k for k in map(str, range(1, 11)) if k not in {str(x) for x in c.get("scorecard", {})}]
        if missing:
            errs.append(f"{n}: scorecard missing criteria {','.join(missing)}")
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("batch")
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--no-check", action="store_true")
    a = ap.parse_args()
    batch = json.load(open(a.batch))
    os.makedirs(a.outdir, exist_ok=True)
    errs = validate(batch)
    if errs:
        print("Fix these in the JSON first:\n  " + "\n  ".join(errs))
        sys.exit(1)

    jobs = [("summary", None)] + [("review", c) for c in batch["candidates"]]
    overflow = []
    for kind, cand in jobs:
        # Try normal size, then shrink fonts slightly before reporting an overflow.
        for scale in (1.0, 0.93, 0.86):
            path = build_summary(batch, a.outdir, scale) if kind == "summary" else build_review(cand, a.outdir, scale)
            pages = None if a.no_check else page_count(path)
            if pages is None or pages <= 1:
                break
        status = "unchecked" if pages is None else f"{pages} page{'s' if pages != 1 else ''}"
        print(f"wrote {path} ({status}{', font scaled %.0f%%' % (scale * 100) if scale < 1 else ''})")
        if pages and pages > 1:
            overflow.append(os.path.basename(path))
    if overflow:
        print("\nOVERFLOW: these still exceed one page; shorten bullets in the JSON and rerun: " + ", ".join(overflow))
        sys.exit(2)


if __name__ == "__main__":
    main()
